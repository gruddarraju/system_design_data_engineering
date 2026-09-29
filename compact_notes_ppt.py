from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).parent
SRC = ROOT / "notes_ppt_updated.docx"
OUT = ROOT / "notes_ppt_compact.docx"

NAVY = RGBColor(12, 25, 48)
TEAL = RGBColor(0, 150, 142)
BLUE = RGBColor(55, 105, 190)
PURPLE = RGBColor(112, 76, 185)
ORANGE = RGBColor(220, 120, 35)
GREEN = RGBColor(45, 135, 70)
DARK = RGBColor(37, 48, 66)
MID = RGBColor(100, 115, 135)
LIGHT_HEX = "DDE6F0"

LABEL_COLORS = {
    "Purpose:": TEAL,
    "How to explain it:": BLUE,
    "Simple example:": PURPLE,
    "Key takeaway:": ORANGE,
    "AI ERP relevance:": GREEN,
    "Companion guide:": NAVY,
}

SPECIAL_PREFIXES = (
    "Company context from the supplied guide:",
    "Candidate profile from the supplied guide:",
    "Interview sequence from the supplied guide:",
    "Preparation reminder:",
    "Campfire attribution and freshness note:",
)


def set_repeatable_styles(doc):
    section = doc.sections[0]
    section.top_margin = Inches(0.42)
    section.bottom_margin = Inches(0.42)
    section.left_margin = Inches(0.52)
    section.right_margin = Inches(0.52)
    section.header_distance = Inches(0.18)
    section.footer_distance = Inches(0.18)

    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(9.2)
    normal.font.color.rgb = DARK
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(2)
    normal.paragraph_format.line_spacing = 1.0
    normal.paragraph_format.widow_control = True

    title = doc.styles["Title"]
    title.font.name = "Aptos Display"
    title.font.size = Pt(24)
    title.font.bold = True
    title.font.color.rgb = NAVY
    title.paragraph_format.space_after = Pt(8)
    title.paragraph_format.keep_with_next = True

    h1 = doc.styles["Heading 1"]
    h1.font.name = "Aptos Display"
    h1.font.size = Pt(15)
    h1.font.bold = True
    h1.font.color.rgb = NAVY
    h1.paragraph_format.space_before = Pt(9)
    h1.paragraph_format.space_after = Pt(3)
    h1.paragraph_format.keep_with_next = True

    h2 = doc.styles["Heading 2"]
    h2.font.name = "Aptos Display"
    h2.font.size = Pt(11.5)
    h2.font.bold = True
    h2.font.color.rgb = TEAL
    h2.paragraph_format.space_before = Pt(5)
    h2.paragraph_format.space_after = Pt(1.5)
    h2.paragraph_format.keep_with_next = True


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = paragraph.add_run("Page ")
    r.font.name = "Aptos"
    r.font.size = Pt(7.5)
    r.font.color.rgb = MID
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    paragraph._p.append(fld)


def add_header_footer(doc):
    section = doc.sections[0]
    hp = section.header.paragraphs[0]
    hp.text = "CAMPFIRE-TAILORED  |  DATA ENGINEERING SYSTEM DESIGN NOTES"
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hp.paragraph_format.space_after = Pt(0)
    for r in hp.runs:
        r.font.name = "Aptos"
        r.font.size = Pt(7.3)
        r.font.bold = True
        r.font.color.rgb = TEAL
    pPr = hp._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "4")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), LIGHT_HEX)
    borders.append(bottom)
    pPr.append(borders)
    add_page_number(section.footer.paragraphs[0])


def add_labeled_paragraph(doc, label, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2.5 if label in ("Key takeaway:", "AI ERP relevance:") else 1.5)
    p.paragraph_format.line_spacing = 1.0
    r = p.add_run(label + " ")
    r.bold = True
    r.font.color.rgb = LABEL_COLORS.get(label, NAVY)
    r = p.add_run(text.strip())
    r.font.color.rgb = DARK
    return p


def add_special_paragraph(doc, source_paragraph):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    for run in source_paragraph.runs:
        r = p.add_run(run.text)
        r.bold = bool(run.bold)
        r.italic = bool(run.italic)
        r.font.color.rgb = NAVY if r.bold else DARK
    if not source_paragraph.runs:
        p.add_run(source_paragraph.text)
    return p


def normalize_join(parts):
    # Imported notes contain hard line wraps as separate paragraphs.
    # Joining on one space preserves words and punctuation while removing layout artifacts.
    return " ".join(x.strip() for x in parts if x.strip())


def main():
    source = Document(SRC)
    compact = Document()
    # Document() may or may not expose an initial empty paragraph, depending on library version.
    if compact.paragraphs and not compact.paragraphs[0].text:
        empty = compact.paragraphs[0]
        empty._element.getparent().remove(empty._element)
    set_repeatable_styles(compact)
    add_header_footer(compact)

    current_label = None
    buffer = []

    def flush():
        nonlocal current_label, buffer
        if buffer:
            text = normalize_join(buffer)
            if current_label:
                add_labeled_paragraph(compact, current_label, text)
            else:
                p = compact.add_paragraph(text)
                p.paragraph_format.space_after = Pt(2)
        current_label = None
        buffer = []

    for index, source_p in enumerate(source.paragraphs):
        text = source_p.text.strip()
        if not text:
            continue

        style = source_p.style.name
        if index == 0 or style == "Title":
            flush()
            p = compact.add_paragraph(style="Title")
            p.add_run(text)
            continue

        if style == "Heading 1":
            flush()
            p = compact.add_heading(text, level=1)
            # Thin divider helps scanning without consuming a blank line.
            pPr = p._p.get_or_add_pPr()
            borders = OxmlElement("w:pBdr")
            bottom = OxmlElement("w:bottom")
            bottom.set(qn("w:val"), "single")
            bottom.set(qn("w:sz"), "5")
            bottom.set(qn("w:space"), "1")
            bottom.set(qn("w:color"), LIGHT_HEX)
            borders.append(bottom)
            pPr.append(borders)
            continue

        if style == "Heading 2":
            flush()
            compact.add_heading(text, level=2)
            continue

        if text in LABEL_COLORS:
            flush()
            current_label = text
            continue

        if text.startswith(SPECIAL_PREFIXES):
            flush()
            add_special_paragraph(compact, source_p)
            continue

        buffer.append(text)

    flush()

    src_props = source.core_properties
    props = compact.core_properties
    props.title = "Compact Campfire-Tailored Data Engineering System Design Notes"
    props.subject = src_props.subject or "Compact senior data engineering and AI ERP interview notes"
    props.author = src_props.author or "Kiro"
    props.keywords = (src_props.keywords or "") + ", compact notes"
    props.comments = "Compact reformat of notes_ppt_updated.docx; content preserved and imported line-wrap paragraphs consolidated."

    compact.save(OUT)
    print(f"Created {OUT}")
    print(f"Bytes: {OUT.stat().st_size}")
    print(f"Source paragraphs: {len([p for p in source.paragraphs if p.text.strip()])}")
    print(f"Compact paragraphs: {len([p for p in compact.paragraphs if p.text.strip()])}")


if __name__ == "__main__":
    main()
